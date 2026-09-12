from pathlib import Path
import json

import numpy as np
import pandas as pd
import pytest

from src.data.imports import import_directory
from tests.fixtures.synthetic import write_sample
from src.database.connection import connect, ROOT
from src.pipeline import load_config, run_portfolio
from src.reporting.daily_report import export_reports
from src.controls.data_quality import quality_checks


@pytest.fixture(scope="module")
def daily_result(tmp_path_factory):
    folder = tmp_path_factory.mktemp("integrated")
    write_sample(folder / "input")
    connection = connect(folder / "test.db")
    import_directory(connection, folder / "input")
    config = load_config("tests/fixtures/portfolio.yaml")
    config["database"] = str(folder / "test.db")
    result = run_portfolio(connection, config)
    export_reports(result, folder / "reports")
    yield connection, result, folder
    connection.close()


def test_end_to_end_reconciliations(daily_result):
    connection, r, folder = daily_result
    assert connection.execute("SELECT COUNT(*) FROM risk_results").fetchone()[0] == 12
    assert abs(r["nav"] - r["positions"].market_value.sum()) < 0.01
    assert r["attribution"]["current_var"] == pytest.approx(
        r["risk"]["summary"]
        .query("method=='Parametric' and confidence==.99")["var"]
        .iloc[0]
    )
    assert abs(r["attribution"]["residual"]) < 1e-7
    assert not (r["quality"].severity == "FAIL").any()
    for _, group in r["risk"]["summary"].groupby("method"):
        assert group.sort_values("confidence")["var"].is_monotonic_increasing
        assert (group.es >= group["var"]).all()
    assert len(r["risk"]["mc_pnl"]) == 10000
    report = json.loads((folder / "reports/risk_report.json").read_text())
    assert report["nav"] == r["nav"]
    assert "synthetic" in report["metadata"]["source_label"]
    assert "95" in (folder / "reports/daily_risk_report.md").read_text()
    assert all(
        abs(row[-1]) < 1e-6
        for row in connection.execute("SELECT * FROM trade_completeness")
    )


def test_post_valuation_controls(daily_result):
    _, r, _ = daily_result
    positions = r["positions"].copy()
    positions.loc[0, "quantity"] += 1
    positions.loc[1, "market_value"] += 100
    positions = positions.iloc[:-1]
    quality = quality_checks(r["tables"], "2026-08-31", positions=positions)
    assert {"POSITION_RECON", "VALUE_RECON"} <= set(
        quality.loc[quality.severity == "FAIL", "check_id"]
    )


def test_failures_persist_and_block_risk(tmp_path):
    write_sample(tmp_path)
    connection = connect(":memory:")
    import_directory(connection, tmp_path)
    connection.execute("DELETE FROM market_prices WHERE security_id='CA_EQ'")
    connection.commit()
    with pytest.raises(ValueError, match="Data quality FAIL"):
        run_portfolio(connection, load_config("tests/fixtures/portfolio.yaml"))
    assert connection.execute("SELECT COUNT(*) FROM risk_results").fetchone()[0] == 0
    assert (
        connection.execute(
            "SELECT COUNT(*) FROM quality_results WHERE severity='FAIL'"
        ).fetchone()[0]
        > 0
    )


def test_dashboard_all_pages(daily_result):
    from streamlit.testing.v1 import AppTest

    _, _, folder = daily_result
    app = AppTest.from_file(str(ROOT / "dashboard/app.py"), default_timeout=30).run()
    app.text_input[0].set_value(str(folder / "reports")).run()
    assert not app.exception
    for page in [
        "Portfolio Overview",
        "Market Risk",
        "Fixed Income",
        "Stress Testing",
        "Risk Attribution",
        "Data Quality",
    ]:
        app.sidebar.radio[0].set_value(page).run()
        assert not app.exception, (page, app.exception)
    app.sidebar.radio[0].set_value("Stress Testing").run()
    app.button[0].click().run()
    assert not app.exception


def test_coupon_receipt_ledger(daily_result):
    _, r, _ = daily_result
    income = r["tables"]["trades"].query("trade_type=='income'").quantity.sum()
    coupon = 0
    for row in r["positions"].query("security_type=='bond'").itertuples():
        coupon += row.quantity * row.face_value * row.coupon_rate / row.coupon_frequency
    assert income == pytest.approx(coupon)


def test_excel_export_refreshes_inputs_and_preserves_risk_snapshots(daily_result):
    from openpyxl import load_workbook
    from src.reporting.excel_export import export_excel

    _, result, folder = daily_result
    source = folder / "reports/risk_report.json"
    data = json.loads(source.read_text())
    row = data["positions"][0]
    change = row["market_price"] * row["multiplier"] * row["fx_rate"]
    row["quantity"] += 1
    row["market_value"] += change
    data["nav"] += change
    updated = folder / "changed_report.json"
    updated.write_text(json.dumps(data))
    book = export_excel(updated, folder / "changed.xlsx")
    values = load_workbook(book, data_only=True)
    formulas = load_workbook(book, data_only=False)
    assert values["Portfolio Summary"]["B6"].value == pytest.approx(
        result["nav"] + change
    )
    assert values["Portfolio Summary"]["B9"].value == pytest.approx(
        result["daily_pnl"] + change
    )
    assert values["Positions"]["H6"].value == row["asset_class"]
    assert all(values["Checks"][f"D{i}"].value == "PASS" for i in range(6, 10))
    assert formulas["Portfolio Summary"]["B6"].value.startswith("=SUM(")
    assert formulas.calculation.fullCalcOnLoad
    assert values["VaR"]["D6"].value == data["risk_summary"][0]["var"]


def test_excel_export_rejects_inconsistent_nav(daily_result):
    from src.reporting.excel_export import export_excel

    _, _, folder = daily_result
    data = json.loads((folder / "reports/risk_report.json").read_text())
    data["nav"] += 100
    source = folder / "bad_report.json"
    source.write_text(json.dumps(data))
    with pytest.raises(ValueError, match="NAV reconciliation"):
        export_excel(source, folder / "bad.xlsx")
