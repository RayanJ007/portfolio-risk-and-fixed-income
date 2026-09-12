from pathlib import Path

import pytest
from openpyxl import load_workbook


def test_delivered_workbook_reconciles_and_preserves_formulas():
    path = Path("models/risk_reporting_model.xlsx")
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    assert len(values.sheetnames) == 16
    assert all(values["Checks"][f"D{r}"].value == "PASS" for r in range(6, 10))
    position_values = [values["Positions"][f"F{r}"].value for r in range(6, 14)]
    assert values["Portfolio Summary"]["B6"].value == pytest.approx(sum(position_values), abs=0.01)
    assert {values["Positions"][f"A{r}"].value for r in range(6, 14)} >= {"SPY", "XIU.TO", "TLT"}
    assert formulas["Positions"]["F6"].value == "=B6*C6*D6*E6"
    lookup = formulas["Positions"]["H6"].value
    assert "XLOOKUP" in (lookup if isinstance(lookup, str) else lookup.text)
    assert len(formulas["Portfolio Summary"]._charts) == 1
    assert formulas["Market Data"].freeze_panes is not None
    errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
    assert not any(
        cell.data_type == "e" and cell.value in errors
        for sheet in values
        for row in sheet
        for cell in row
    )
